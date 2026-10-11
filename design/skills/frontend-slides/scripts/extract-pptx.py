#!/usr/bin/env python3
"""
Extract all content from a PowerPoint file (.pptx).
Returns a JSON structure with slides, text, and images.

Usage:
    python extract-pptx.py <input.pptx> [output_dir]

Text includes tables and shapes inside groups; images include pictures placed
in picture placeholders. Writes <output_dir>/extracted-slides.json and
<output_dir>/assets/.

Requires: pip install python-pptx
"""

import json
import os
import sys

try:
    from pptx import Presentation
    from pptx.shapes.picture import Picture
except ImportError:
    print("Error: python-pptx is required: pip install python-pptx", file=sys.stderr)
    sys.exit(1)


def iter_shapes(shapes):
    """Yield shapes, descending into group shapes (their children are otherwise skipped)."""
    for shape in shapes:
        if hasattr(shape, "shapes"):  # GroupShape
            yield from iter_shapes(shape.shapes)
        else:
            yield shape


def extract_pptx(file_path, output_dir="."):
    """
    Extract all content from a PowerPoint file.
    Returns a list of slide data dicts with text, images, and notes.
    """
    prs = Presentation(file_path)
    slides_data = []

    # Create assets directory for extracted images
    assets_dir = os.path.join(output_dir, "assets")
    os.makedirs(assets_dir, exist_ok=True)

    for slide_num, slide in enumerate(prs.slides):
        slide_data = {
            "number": slide_num + 1,
            "title": "",
            "content": [],
            "images": [],
            "notes": "",
        }

        for shape in iter_shapes(slide.shapes):
            # Extract table content (cell text, row by row)
            if getattr(shape, "has_table", False):
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                slide_data["content"].append({"type": "table", "rows": rows})

            # Extract text content
            if shape.has_text_frame:
                if shape == slide.shapes.title:
                    slide_data["title"] = shape.text
                else:
                    slide_data["content"].append(
                        {"type": "text", "content": shape.text}
                    )

            # Extract images (plain pictures and pictures in placeholders;
            # linked, non-embedded pictures have no blob and are skipped)
            if isinstance(shape, Picture):
                try:
                    image = shape.image
                except (KeyError, ValueError, AttributeError):
                    continue
                image_bytes = image.blob
                image_ext = image.ext
                image_name = f"slide{slide_num + 1}_img{len(slide_data['images']) + 1}.{image_ext}"
                image_path = os.path.join(assets_dir, image_name)

                with open(image_path, "wb") as f:
                    f.write(image_bytes)

                slide_data["images"].append(
                    {
                        "path": f"assets/{image_name}",
                        "width": shape.width,
                        "height": shape.height,
                    }
                )

        # Extract speaker notes
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            slide_data["notes"] = notes_frame.text if notes_frame is not None else ""

        slides_data.append(slide_data)

    return slides_data


if __name__ == "__main__":
    usage = "Usage: python extract-pptx.py <input.pptx> [output_dir]"
    if len(sys.argv) > 1 and sys.argv[1] in ("-h", "--help"):
        print(usage)
        sys.exit(0)
    if len(sys.argv) < 2 or len(sys.argv) > 3:
        print(usage, file=sys.stderr)
        sys.exit(1)

    input_file = sys.argv[1]
    output_dir = sys.argv[2] if len(sys.argv) > 2 else "."
    if not os.path.isfile(input_file):
        print(f"Error: file not found: {input_file}", file=sys.stderr)
        sys.exit(1)

    try:
        slides = extract_pptx(input_file, output_dir)
    except Exception as exc:  # e.g. not a zip / not a presentation package
        print(f"Error: could not extract {input_file} (is it a .pptx file?): "
              f"{exc.__class__.__name__}: {exc}", file=sys.stderr)
        sys.exit(1)

    # Write extracted data as JSON
    output_path = os.path.join(output_dir, "extracted-slides.json")
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(slides, f, indent=2)

    print(f"Extracted {len(slides)} slides to {output_path}")
    for s in slides:
        img_count = len(s["images"])
        print(f"  Slide {s['number']}: {s['title'] or '(no title)'} — {img_count} image(s)")
